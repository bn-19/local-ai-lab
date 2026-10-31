"""Convert UCI Online Retail workbook without altering/deduplicating rows.
Source https://archive.ics.uci.edu/dataset/352/online+retail (CC BY 4.0).
InvoiceDate is emitted ISO 8601; all other source values are preserved.
"""
import csv, datetime, hashlib, json, time
from pathlib import Path
from openpyxl import load_workbook
source=Path('.local-lab/sources/retail/Online Retail.xlsx')
out=source.with_suffix('.csv')
start=time.perf_counter()
wb=load_workbook(source,read_only=True,data_only=True)
count=0; net=0; positive=0
from decimal import Decimal
net=Decimal(0); positive=Decimal(0)
with out.open('w',newline='') as f:
 writer=csv.writer(f)
 for i,row in enumerate(wb.active.iter_rows(values_only=True)):
  writer.writerow([v.isoformat(sep=' ') if isinstance(v,datetime.datetime) else v for v in row])
  if i:
   count+=1
   q,p=Decimal(str(row[3])),Decimal(str(row[5]))
   net+=q*p
   if q>=0 and p>=0:positive+=q*p
wb.close()
evidence={'source':str(source),'csv':str(out),'rows':count,'bytes':out.stat().st_size,'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'net_revenue':str(net),'nonnegative_revenue':str(positive),'conversionSeconds':round(time.perf_counter()-start,3),'transformation':'All source rows, dates ISO8601; no duplicated/generated records','license':'CC BY 4.0','attribution':'Daqing Chen, UCI Online Retail, DOI 10.24432/C5BW33'}
Path('.local-lab/retail-provenance.json').write_text(json.dumps(evidence,indent=2))
print(json.dumps(evidence,indent=2))
